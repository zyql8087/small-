"""Generate a 4-slide group meeting PPT based on updated docx report (Miura-ori + TMP dataset)."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Color scheme ---
BG_COLOR = RGBColor(0xF5, 0xF5, 0xF5)
ACCENT = RGBColor(0x1A, 0x56, 0x8E)
ACCENT2 = RGBColor(0x2E, 0x86, 0xAB)
TEXT_DARK = RGBColor(0x2D, 0x2D, 0x2D)
TEXT_GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HEADER_BG = RGBColor(0x1A, 0x56, 0x8E)
TABLE_ROW_ALT = RGBColor(0xE8, 0xF0, 0xF7)
HIGHLIGHT = RGBColor(0xE8, 0x6D, 0x3F)
GREEN = RGBColor(0x27, 0xAE, 0x60)

# --- Image paths ---
IMG_CANOPY = r"F:\small++\external\GenerateOrigamiDataSet\Figures_ReadMe\CanopyDetails.png"
IMG_FWD_LOSS = r"F:\small++\origami_experiments\model_audit\loss_curves_final\forward_transformer_loss.png"
IMG_FWD_NRMSE = r"F:\small++\origami_experiments\results_final_model_test\figures\forward_overall_nrmse.png"
IMG_INV_NRMSE = r"F:\small++\origami_experiments\results_final_model_test\figures\inverse_closed_loop_nrmse.png"
IMG_SHOWCASE = r"F:\small++\origami_experiments\results_final_model_test\prediction_showcase\forward_transformer_good_case_1_sample_165.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape_rect(slide, left, top, width, height, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=14,
                bold=False, color=TEXT_DARK, alignment=PP_ALIGN.LEFT, font_name="微软雅黑"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=13, color=TEXT_DARK):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "微软雅黑"
        p.space_after = Pt(6)
    return txBox


def add_card(slide, left, top, width, height, title, items, title_color=ACCENT):
    add_shape_rect(slide, left, top, width, height, WHITE, RGBColor(0xDD, 0xDD, 0xDD))
    bar = add_shape_rect(slide, left, top, width, Inches(0.45), title_color)
    bar.line.fill.background()
    add_textbox(slide, left + Inches(0.15), top + Inches(0.05), width - Inches(0.3), Inches(0.4),
                title, font_size=14, bold=True, color=WHITE)
    add_bullet_list(slide, left + Inches(0.2), top + Inches(0.55),
                    width - Inches(0.4), height - Inches(0.65), items, font_size=12)


def add_table(slide, left, top, width, height, headers, rows, col_widths=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = WHITE
            p.font.name = "微软雅黑"; p.alignment = PP_ALIGN.CENTER
        cell.fill.solid(); cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11); p.font.color.rgb = TEXT_DARK
                p.font.name = "微软雅黑"; p.alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = TABLE_ROW_ALT if r % 2 == 1 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return table_shape


def add_image(slide, img_path, left, top, width=None, height=None):
    if not os.path.exists(img_path):
        print(f"  WARNING: Image not found: {img_path}")
        return None
    if width and height:
        return slide.shapes.add_picture(img_path, left, top, width, height)
    elif width:
        return slide.shapes.add_picture(img_path, left, top, width=width)
    elif height:
        return slide.shapes.add_picture(img_path, left, top, height=height)
    return slide.shapes.add_picture(img_path, left, top)


def add_slide_number(slide, num, total=4):
    add_textbox(slide, SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.5),
                Inches(1), Inches(0.4), f"{num} / {total}",
                font_size=10, color=TEXT_GRAY, alignment=PP_ALIGN.RIGHT)


# ================================================================
# SLIDE 1: Dataset Introduction (Miura-ori + TMP)
# ================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide1, BG_COLOR)

bar1 = add_shape_rect(slide1, Inches(0), Inches(0), SLIDE_W, Inches(1.1), ACCENT)
bar1.line.fill.background()
add_textbox(slide1, Inches(0.6), Inches(0.25), Inches(10), Inches(0.6),
            "Miura-ori 折纸结构数据集介绍", font_size=28, bold=True, color=WHITE)
add_textbox(slide1, Inches(0.6), Inches(1.25), Inches(12), Inches(0.4),
            "折纸结构刚度预测与逆向设计 · 组会汇报", font_size=15, color=TEXT_GRAY)

# Left: two cards
add_card(slide1, Inches(0.4), Inches(1.8), Inches(6.0), Inches(2.6),
         "数据集概览", [
             "数据集：Origami Sheet MaterialProperty（Miura + TMP）",
             "原始来源：MiuraSheetMat.txt + TMPSheetMat.txt，各 2000 条",
             "总样本：4000 条（训练 3200 / 测试 800）",
             "训练集内部切分：train/val = 2560/640",
             "预处理：log10(stiffness) → StandardScaler 标准化",
         ], title_color=ACCENT)

add_card(slide1, Inches(0.4), Inches(4.7), Inches(6.0), Inches(2.5),
         "任务定义", [
             "正向任务：8 个设计/材料参数 → 6 个刚度指标",
             "逆向任务：6 个目标刚度 → 反推 8 个设计参数",
             "逆向评估：参数恢复 + 闭环刚度匹配（surrogate 验证）",
             "对比模型：ResMLP、GAT、Transformer、CVAE、Diffusion",
         ], title_color=ACCENT2)

# Right: dataset structure table + image
add_textbox(slide1, Inches(6.8), Inches(1.85), Inches(6), Inches(0.35),
            "输入参数（8个）", font_size=13, bold=True, color=ACCENT)

inp_headers = ["参数", "含义", "范围"]
inp_rows = [
    ["pattern", "折纸图案类型", "1=Miura, 2=TMP"],
    ["m / n", "单元格数量", "24,30,36 / 6,9,12"],
    ["tcrease", "折痕厚度", "0.5e-3~1.0e-3 m"],
    ["tpanel", "面板厚度", "1e-3~6e-3 m"],
    ["W", "折痕宽度", "1e-3~4e-3 m"],
    ["creaseE", "折痕杨氏模量", "1e9~5e9 Pa"],
    ["panelE", "面板杨氏模量", "1e9~5e9 Pa"],
]
add_table(slide1, Inches(6.8), Inches(2.25), Inches(6.1), Inches(2.5),
          inp_headers, inp_rows, col_widths=[Inches(1.5), Inches(2.0), Inches(2.6)])

add_textbox(slide1, Inches(6.8), Inches(4.95), Inches(6), Inches(0.35),
            "输出目标（6个刚度值）", font_size=13, bold=True, color=ACCENT2)

out_headers = ["刚度指标", "含义"]
out_rows = [
    ["bendstiff30/60/90", "30%/60%/90% 部署水平下弯曲刚度"],
    ["axialstiff30/60/90", "30%/60%/90% 部署水平下轴向刚度"],
]
add_table(slide1, Inches(6.8), Inches(5.35), Inches(6.1), Inches(1.0),
          out_headers, out_rows, col_widths=[Inches(2.5), Inches(3.6)])

# Canopy image as small thumbnail
add_image(slide1, IMG_CANOPY,
          Inches(6.8), Inches(6.5), height=Inches(0.7))

add_slide_number(slide1, 1)

# ================================================================
# SLIDE 2: Forward Models - Training + Results
# ================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide2, BG_COLOR)

bar2 = add_shape_rect(slide2, Inches(0), Inches(0), SLIDE_W, Inches(1.1), ACCENT)
bar2.line.fill.background()
add_textbox(slide2, Inches(0.6), Inches(0.25), Inches(10), Inches(0.6),
            "正向模型：训练与测试集性能", font_size=28, bold=True, color=WHITE)

# Left: training curves image
add_image(slide2, IMG_FWD_LOSS,
          Inches(0.4), Inches(1.3), width=Inches(6.3))

add_textbox(slide2, Inches(0.4), Inches(5.5), Inches(6.3), Inches(0.3),
            "▲ Forward Transformer 训练/验证 Loss 曲线",
            font_size=10, color=TEXT_GRAY, alignment=PP_ALIGN.CENTER)

# Training config mini-card below loss curve
add_card(slide2, Inches(0.4), Inches(5.9), Inches(6.3), Inches(1.3),
         "训练要点", [
             "所有模型使用 MSE Loss，Adam/AdamW 优化器",
             "训练集内 80/20 切分验证，测试集仅最终评估",
             "6 个刚度目标统一 log10 + 标准化后训练",
         ], title_color=ACCENT)

# Right: NRMSE chart + results table
add_image(slide2, IMG_FWD_NRMSE,
          Inches(6.9), Inches(1.3), width=Inches(6.0))

add_textbox(slide2, Inches(6.9), Inches(3.9), Inches(6.0), Inches(0.3),
            "▲ 正向模型整体 NRMSE% 对比（Train / Val / Test）",
            font_size=10, color=TEXT_GRAY, alignment=PP_ALIGN.CENTER)

# Results summary
add_card(slide2, Inches(6.9), Inches(4.3), Inches(6.0), Inches(3.0),
         "正向模型测试集结论", [
             "三个正向模型测试 R² 均超过 0.999",
             "ResMLP 整体 NRMSE 最低（~0.16%）",
             "Transformer MAE 最低",
             "GAT 稍弱，可作为物理图建模 baseline",
             "5 pipelines 跑 3 seeds，结果稳定可靠",
         ], title_color=GREEN)

add_slide_number(slide2, 2)

# ================================================================
# SLIDE 3: Inverse Models + Prediction Showcase
# ================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide3, BG_COLOR)

bar3 = add_shape_rect(slide3, Inches(0), Inches(0), SLIDE_W, Inches(1.1), ACCENT)
bar3.line.fill.background()
add_textbox(slide3, Inches(0.6), Inches(0.25), Inches(10), Inches(0.6),
            "逆向模型：闭环匹配与预测样例", font_size=28, bold=True, color=WHITE)

# Left: inverse NRMSE chart
add_image(slide3, IMG_INV_NRMSE,
          Inches(0.4), Inches(1.3), width=Inches(6.0))

add_textbox(slide3, Inches(0.4), Inches(4.0), Inches(6.0), Inches(0.3),
            "▲ 逆向模型闭环刚度 NRMSE% 对比",
            font_size=10, color=TEXT_GRAY, alignment=PP_ALIGN.CENTER)

# Left bottom: inverse findings
add_card(slide3, Inches(0.4), Inches(4.4), Inches(6.0), Inches(2.8),
         "逆向模型结论", [
             "Inverse ResMLP：连续参数 R² 最高，适合参数恢复",
             "Inverse Diffusion x0：闭环 NRMSE 最低，适合目标刚度设计",
             "CVAE：闭环优于普通 ResMLP，但参数恢复弱",
             "  → 生成式模型可产生多解，不一定复现原始参数",
             "split diffusion < mixed diffusion：低维表格数据不必完全拆分离散/连续",
         ], title_color=ACCENT2)

# Right: prediction showcase
add_image(slide3, IMG_SHOWCASE,
          Inches(6.7), Inches(1.3), width=Inches(6.2))

add_textbox(slide3, Inches(6.7), Inches(5.2), Inches(6.2), Inches(0.3),
            "▲ Forward Transformer 预测样例（平均相对误差 ~0.238%）",
            font_size=10, color=TEXT_GRAY, alignment=PP_ALIGN.CENTER)

# Right bottom: key insight
add_card(slide3, Inches(6.7), Inches(5.6), Inches(6.2), Inches(1.6),
         "关键洞察", [
             "原始刚度 MAE 绝对值大 → 刚度量级差异大（10²~10⁶）",
             "NRMSE 和 R² 才能反映真实的相对预测质量",
             "正向预测精度极高（R²>0.999），已可用于实际设计筛选",
         ], title_color=GREEN)

add_slide_number(slide3, 3)

# ================================================================
# SLIDE 4: Conclusions & Recommendations
# ================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide4, BG_COLOR)

bar4 = add_shape_rect(slide4, Inches(0), Inches(0), SLIDE_W, Inches(1.1), ACCENT)
bar4.line.fill.background()
add_textbox(slide4, Inches(0.6), Inches(0.25), Inches(10), Inches(0.6),
            "结论与模型使用建议", font_size=28, bold=True, color=WHITE)

# Conclusion card
add_shape_rect(slide4, Inches(0.4), Inches(1.3), Inches(12.4), Inches(2.6),
               WHITE, RGBColor(0xDD, 0xDD, 0xDD))
cbar = add_shape_rect(slide4, Inches(0.4), Inches(1.3), Inches(12.4), Inches(0.5), ACCENT)
cbar.line.fill.background()
add_textbox(slide4, Inches(0.55), Inches(1.35), Inches(11), Inches(0.4),
            "主要结论", font_size=16, bold=True, color=WHITE)

conclusions = [
    "1. 正向预测精度极高（R²>0.999），ResMLP NRMSE 最低，Transformer MAE 最低，GAT 可作图建模 baseline",
    "2. 逆向设计需区分「参数恢复」与「闭环刚度匹配」：ResMLP 参数恢复最优，Diffusion x0 闭环匹配最优",
    "3. CVAE 闭环优于 ResMLP 但参数恢复弱——生成式模型可多解但不一定复现原始参数",
    "4. 原始刚度 MAE 绝对值大因量级差异（10²~10⁶），log10+标准化后 NRMSE/R² 才能反映真实预测质量",
]
for i, c in enumerate(conclusions):
    add_textbox(slide4, Inches(0.7), Inches(2.0 + i * 0.45), Inches(12), Inches(0.4),
                c, font_size=13, color=TEXT_DARK)

# Recommendation table
rec_headers = ["使用场景", "推荐模型", "理由"]
rec_rows = [
    ["正向刚度预测", "ResMLP / Transformer", "R²>0.999，NRMSE<0.2%，工业级精度"],
    ["逆向参数恢复", "Inverse ResMLP", "连续参数 R² 最高"],
    ["逆向目标刚度设计", "Inverse Diffusion x0", "闭环 NRMSE 最低，适合工程应用"],
    ["快速 baseline", "GAT", "物理图建模，可解释性强"],
    ["多候选方案生成", "CVAE", "可采样多组参数，探索设计空间"],
]
add_table(slide4, Inches(0.4), Inches(4.2), Inches(12.4), Inches(2.2),
          rec_headers, rec_rows,
          col_widths=[Inches(2.5), Inches(3.0), Inches(6.9)])

# Future work
add_textbox(slide4, Inches(0.4), Inches(6.6), Inches(12.4), Inches(0.5),
            "后续方向：逆向多解评估 → 候选设计可制造性约束 → 连续应力-应变曲线数据表示扩展",
            font_size=12, color=TEXT_GRAY)

add_slide_number(slide4, 4)

# Save
output_path = r"F:\small++\origami_experiments\reports\组会汇报_v3_final.pptx"
prs.save(output_path)
print(f"PPT saved to: {output_path}")
