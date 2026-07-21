"""Generate a DOCX report for the Origami taskfit experiments."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


EXPERIMENT_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_DIR = EXPERIMENT_ROOT
RESULT_DIR = SCRIPT_DIR / "results_final_model_test"
LOSS_DIR = SCRIPT_DIR / "model_audit" / "loss_curves_final"
SHOWCASE_DIR = RESULT_DIR / "prediction_showcase"
REPORT_DIR = SCRIPT_DIR / "reports"
REPORT_PATH = REPORT_DIR / "origami_model_experiment_report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def set_document_style(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        if style_name == "Title":
            style.font.color.rgb = RGBColor(31, 78, 121)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_picture(document: Document, path: Path, caption: str, width: float = 6.3) -> None:
    if not path.exists():
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[缺失图片: {path}]").italic = True
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(document, caption)


def add_table_from_rows(document: Document, rows: list[dict], columns: list[str], title: str | None = None) -> None:
    if title:
        document.add_paragraph(title, style="Heading 3")
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for idx, col in enumerate(columns):
        set_cell_text(header[idx], col, bold=True)
        set_cell_shading(header[idx], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            value = row.get(col, "")
            if isinstance(value, float):
                if abs(value) >= 100000:
                    text = f"{value:.3e}"
                elif abs(value) >= 10:
                    text = f"{value:.2f}"
                else:
                    text = f"{value:.4f}"
            else:
                text = str(value)
            set_cell_text(cells[idx], text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def read_overall_metrics() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    forward = pd.read_csv(RESULT_DIR / "forward_test_metrics.csv")
    inverse_param = pd.read_csv(RESULT_DIR / "inverse_parameter_test_metrics.csv")
    inverse_loop = pd.read_csv(RESULT_DIR / "inverse_closed_loop_test_metrics.csv")
    classifier = pd.read_csv(RESULT_DIR / "inverse_classifier_test_metrics.csv")
    return forward, inverse_param, inverse_loop, classifier


def compact_metric_rows(df: pd.DataFrame, group: str, cols: list[str]) -> list[dict]:
    out = df[df["Group"] == group].copy()
    return out[cols].to_dict("records")


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    forward, inverse_param, inverse_loop, classifier = read_overall_metrics()

    document = Document()
    set_document_style(document)
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("折纸结构刚度预测与逆向设计模型实验报告")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("基于 Origami 数据集的 ResMLP、GAT、Transformer、CVAE 与 Diffusion 对比实验").italic = True

    document.add_paragraph()
    document.add_heading("1. 实验目的", level=1)
    document.add_paragraph(
        "本实验围绕折纸结构的力学刚度预测与逆向设计展开。正向任务是由设计/材料参数预测 6 个刚度指标；"
        "逆向任务是由目标刚度反推结构设计参数，并通过正向 surrogate 检查反推参数能否匹配目标刚度。"
    )
    add_bullets(
        document,
        [
            "建立适配该数据集的正向预测模型，比较 ResMLP、GAT 和 Transformer 的预测性能。",
            "建立逆向模型，比较确定性回归、CVAE 和 Diffusion 在参数恢复与闭环刚度匹配上的表现。",
            "保留所有旧实验结果，在新目录中输出 checkpoint、loss 曲线、测试指标和展示图片。",
        ],
    )

    document.add_heading("2. 数据集与任务定义", level=1)
    document.add_paragraph(
        "数据来自 Origami 折纸结构数据集，包含 Miura 与 TMP 两类折纸模式。整理后的训练文件为 "
        "origami_train.npz，测试文件为 origami_test.npz。总数据量为 4000 条，其中训练集 3200 条，"
        "测试集 800 条；训练过程中训练集内部再固定切分为 train/val = 2560/640。"
    )
    add_table_from_rows(
        document,
        [
            {"类别": "输入参数", "变量": "pattern, m, n", "说明": "离散设计变量，使用分类 id 与 embedding 表示"},
            {"类别": "输入参数", "变量": "tcrease, tpanel, W, creaseE, panelE", "说明": "连续几何/材料变量，使用 log10 + StandardScaler"},
            {"类别": "输出刚度", "变量": "bendstiff30/60/90", "说明": "弯曲刚度 3 个载荷点"},
            {"类别": "输出刚度", "变量": "axialstiff30/60/90", "说明": "轴向刚度 3 个载荷点"},
        ],
        ["类别", "变量", "说明"],
        "表 1 数据字段说明",
    )
    document.add_paragraph(
        "这里 30/60/90 表示 30%、60%、90% 变形/拉伸水平下的刚度采样点，不是角度。"
        "刚度目标统一使用 log10(stiffness) 后标准化，避免数量级较大的刚度项主导损失。"
    )

    document.add_heading("3. 模型与训练策略", level=1)
    add_table_from_rows(
        document,
        [
            {"模型": "Forward ResMLP", "任务": "8 参数 -> 6 刚度", "特点": "离散 embedding + 连续 MLP，弯曲/轴向双输出头"},
            {"模型": "Forward GAT", "任务": "8 参数 -> 6 刚度", "特点": "采用 physical_sparse 物理分组边，作为图模型 baseline"},
            {"模型": "Forward Transformer", "任务": "8 参数 -> 6 刚度", "特点": "将参数视作 token 建模，验证集 loss 最低"},
            {"模型": "Inverse ResMLP", "任务": "6 刚度 -> 8 参数", "特点": "离散分类头 + 连续回归头，适合参数恢复"},
            {"模型": "Inverse CVAE", "任务": "6 刚度 -> 参数分布", "特点": "生成式逆向模型，使用闭环刚度匹配筛选"},
            {"模型": "Inverse Diffusion x0", "任务": "6 刚度 -> 参数分布", "特点": "直接预测 x0，best-of-50 闭环选择，闭环测试最优"},
        ],
        ["模型", "任务", "特点"],
        "表 2 模型结构与用途",
    )
    document.add_paragraph(
        "训练过程中只使用训练集及其内部验证集进行调参；测试集仅在最终性能评估阶段读取。"
        "所有最终模型已导出为 portable checkpoint，避免早期 checkpoint 中 TaskfitPreprocessor pickle 导致的跨脚本加载问题。"
    )

    document.add_heading("4. 训练 Loss 曲线", level=1)
    document.add_paragraph("下图展示代表性模型的训练/验证 loss 曲线。")
    add_picture(document, LOSS_DIR / "forward_transformer_loss.png", "图 1 Forward Transformer 训练与验证 loss", width=6.1)
    add_picture(document, LOSS_DIR / "inverse_cvae_physical_loss.png", "图 2 Inverse CVAE 训练与验证 loss", width=6.1)
    add_picture(document, LOSS_DIR / "inverse_diffusion_x0_loss.png", "图 3 Inverse Diffusion x0 训练与验证 loss", width=6.1)

    document.add_heading("5. 测试集性能对比", level=1)
    document.add_paragraph("最终测试使用 held-out 测试集 origami_test.npz，共 800 条样本。")

    forward_rows = compact_metric_rows(forward, "All", ["Model", "MAE", "RMSE", "R2", "NRMSE%"])
    add_table_from_rows(document, forward_rows, ["Model", "MAE", "RMSE", "R2", "NRMSE%"], "表 3 正向模型整体测试指标")
    add_picture(document, RESULT_DIR / "figures" / "forward_overall_nrmse.png", "图 4 正向模型整体 NRMSE 对比", width=5.8)
    add_picture(document, RESULT_DIR / "figures" / "forward_per_stiffness_mae.png", "图 5 正向模型各刚度分量 MAE 对比", width=6.2)

    loop_rows = compact_metric_rows(inverse_loop, "All", ["Model", "MAE", "RMSE", "R2", "NRMSE%"])
    add_table_from_rows(document, loop_rows, ["Model", "MAE", "RMSE", "R2", "NRMSE%"], "表 4 逆向模型闭环刚度匹配测试指标")
    add_picture(document, RESULT_DIR / "figures" / "inverse_closed_loop_nrmse.png", "图 6 逆向模型闭环刚度 NRMSE 对比", width=6.0)

    param_rows = compact_metric_rows(inverse_param, "continuous_params_all", ["Model", "MAE", "RMSE", "R2", "NRMSE%"])
    add_table_from_rows(document, param_rows, ["Model", "MAE", "RMSE", "R2", "NRMSE%"], "表 5 逆向模型连续参数恢复测试指标")
    add_picture(document, RESULT_DIR / "figures" / "inverse_continuous_param_nrmse.png", "图 7 逆向连续参数恢复 NRMSE 对比", width=5.8)

    add_table_from_rows(
        document,
        classifier[["Model", "Group", "Accuracy"]].to_dict("records"),
        ["Model", "Group", "Accuracy"],
        "表 6 inverse_classifier 离散变量测试准确率",
    )
    add_picture(document, RESULT_DIR / "figures" / "inverse_classifier_accuracy.png", "图 8 离散变量分类准确率", width=5.6)

    document.add_heading("6. 预测样例展示", level=1)
    document.add_paragraph("为了展示模型在具体样本上的预测效果，从测试集中选取了平均相对误差较低的样本绘制结果。")
    add_picture(
        document,
        SHOWCASE_DIR / "forward_transformer_good_case_1_sample_165.png",
        "图 9 正向 Transformer 预测样例，平均相对误差约 0.238%",
        width=6.3,
    )
    add_picture(
        document,
        SHOWCASE_DIR / "forward_resmlp_good_case_1_sample_408.png",
        "图 10 正向 ResMLP 预测样例，平均相对误差约 0.279%",
        width=6.3,
    )
    add_picture(
        document,
        SHOWCASE_DIR / "inverse_diffusion_good_case_1_sample_143.png",
        "图 11 Inverse Diffusion x0 闭环逆向设计样例，平均相对误差约 0.235%",
        width=6.3,
    )

    document.add_heading("7. 结果分析", level=1)
    add_bullets(
        document,
        [
            "正向预测整体精度很高，三个正向模型测试 R2 均超过 0.999；ResMLP 的整体 NRMSE 最低，Transformer 的 MAE 最低，GAT 稍弱但仍可作为物理图建模 baseline。",
            "逆向任务需要区分“参数恢复”和“闭环刚度匹配”。如果目标是恢复原始参数，Inverse ResMLP 的连续参数 R2 最高；如果目标是找到能实现目标刚度的设计，Inverse Diffusion x0 的闭环 NRMSE 最低。",
            "CVAE 在闭环设计上优于普通 Inverse ResMLP，但参数恢复能力弱于 ResMLP，符合生成式模型可产生多解但不一定复现原始参数的特点。",
            "split continuous diffusion 经过 teacher forcing 和 x0 预测优化后仍弱于 mixed diffusion，说明在该低维表格数据集上，将离散与连续目标完全拆开并不一定更优。",
            "原始刚度 MAE 数值看起来较大，主要是由于刚度绝对量级很大；NRMSE 和 R2 更能反映相对预测质量。",
        ],
    )

    document.add_heading("8. 结论与推荐使用方式", level=1)
    add_table_from_rows(
        document,
        [
            {"使用目标": "正向刚度预测", "推荐模型": "Forward ResMLP / Forward Transformer", "原因": "测试集 R2 > 0.999，NRMSE 约 0.25%-0.27%"},
            {"使用目标": "逆向闭环设计", "推荐模型": "Inverse Diffusion x0", "原因": "闭环刚度 NRMSE 最低，适合寻找满足目标刚度的设计"},
            {"使用目标": "真实参数恢复", "推荐模型": "Inverse ResMLP", "原因": "连续参数恢复 R2 最高"},
            {"使用目标": "离散变量判断", "推荐模型": "inverse_classifier", "原因": "pattern/m/n 测试准确率分别为 1.000/0.951/0.995"},
        ],
        ["使用目标", "推荐模型", "原因"],
        "表 7 最终模型使用建议",
    )
    document.add_paragraph(
        "综合来看，本实验已经形成可用的正向预测与逆向设计模型体系。后续若继续提升性能，建议优先围绕逆向多解评估、"
        "候选设计可制造性约束、以及更接近连续应力-应变曲线的数据表示进行扩展。"
    )

    document.save(REPORT_PATH)
    print(f"[Saved] {REPORT_PATH}")


if __name__ == "__main__":
    main()
