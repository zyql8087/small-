"""Create a Chinese manuscript-style DOCX for the inverse SWOMPS validation study."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SCRIPT_DIR = Path(__file__).resolve().parent
EXPERIMENT_DIR = SCRIPT_DIR.parent
WORKSPACE = EXPERIMENT_DIR.parents[2]
OUTPUT_DIR = WORKSPACE / "output" / "doc"
OUTPUT_DOCX = OUTPUT_DIR / "miura_inverse_swomps_validation_draft.docx"

FORWARD_METRICS = EXPERIMENT_DIR / "training_results" / "curve_forward_quality_filtered_random_50ep_v13" / "metrics.json"
INVERSE_METRICS = (
    EXPERIMENT_DIR
    / "training_results"
    / "inverse_diffusion_quality_filtered_150ep_v2_physical"
    / "inverse_diffusion_closed_loop_metrics.json"
)
VALIDATION_DIR = EXPERIMENT_DIR / "validation" / "inverse_swomps_top10_v1"
SWOMPS_SUMMARY = VALIDATION_DIR / "results" / "swomps_validation_summary.json"
SWOMPS_PER_CANDIDATE = VALIDATION_DIR / "results" / "swomps_validation_metrics.csv"
AUDIT_SUMMARY = VALIDATION_DIR / "audit_nonlinear" / "miura_curve_audit.json"

FIG_FORWARD = EXPERIMENT_DIR / "training_results" / "curve_forward_quality_filtered_random_50ep_v13" / "training_loss_curves.png"
FIG_INVERSE = (
    EXPERIMENT_DIR
    / "training_results"
    / "inverse_diffusion_quality_filtered_150ep_v2_physical"
    / "closed_loop_nrmse.png"
)
FIG_SWOMPS_SUMMARY = VALIDATION_DIR / "results" / "swomps_validation_nrmse.png"
FIG_SWOMPS_GOOD = VALIDATION_DIR / "results" / "curve_plots" / "invval_0001_miura_00074_true_swomps.png"
FIG_SWOMPS_WORST = VALIDATION_DIR / "results" / "curve_plots" / "invval_0007_miura_00968_true_swomps.png"
FIG_PARAMETER_SCHEMATIC = OUTPUT_DIR / "miura_parameter_mapping_schematic.png"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(9)


def set_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for name, size in [("Title", 18), ("Heading 1", 14), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_para(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.first_line_indent = Pt(21) if style is None else None
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9)
    run.italic = True


def add_figure(document: Document, path: Path, caption: str, width: float = 5.8) -> None:
    if not path.exists():
        add_para(document, f"[图像缺失: {path}]")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(document, caption)


def add_table(document: Document, headers: list[str], rows: list[list[str]], title: str) -> None:
    add_caption(document, title)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    document.add_paragraph()


def create_parameter_schematic() -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Circle, FancyArrowPatch, Polygon, Rectangle

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(13.8, 7.4), dpi=220)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 9)
    ax.set_aspect("equal")
    ax.axis("off")

    blue = "#0b57ff"
    red = "#e00000"
    crease_red = "#c0392b"
    panel_fill = "#d9d9d9"

    ax.text(
        8,
        8.55,
        "CURRENT Miura-ori Simulation Loading Process (Not Top Platen Compression)",
        ha="center",
        va="center",
        fontsize=17,
        weight="bold",
    )

    def draw_miura_sheet(x0: float, y0: float, scale: float = 1.0) -> tuple[list[tuple[float, float]], list[tuple[float, float]]]:
        rows, cols = 4, 8
        amp = 0.34

        def vertex(i: int, j: int) -> tuple[float, float, float]:
            x = i * 0.72 + 0.18 * (j % 2)
            y = j * 0.58 + 0.08 * ((-1) ** i)
            z = amp * ((-1) ** i) * (0.95 - 0.05 * (j % 2))
            return x, y, z

        def project(point: tuple[float, float, float]) -> tuple[float, float]:
            x, y, z = point
            return (
                x0 + scale * (0.93 * x + 0.44 * y),
                y0 + scale * (-0.13 * x + 0.43 * y + 0.82 * z),
            )

        grid3 = [[vertex(i, j) for j in range(rows + 1)] for i in range(cols + 1)]
        grid2 = [[project(grid3[i][j]) for j in range(rows + 1)] for i in range(cols + 1)]
        all_points = [grid2[i][j] for i in range(cols + 1) for j in range(rows + 1)]
        min_x = min(p[0] for p in all_points)
        max_x = max(p[0] for p in all_points)
        min_y = min(p[1] for p in all_points)
        max_y = max(p[1] for p in all_points)
        shadow = Polygon(
            [(min_x + 0.10, min_y - 0.18), (max_x + 0.24, min_y - 0.12), (max_x - 0.15, min_y + 0.12), (min_x - 0.12, min_y + 0.04)],
            closed=True,
            facecolor="#b8b8b8",
            edgecolor="none",
            alpha=0.22,
            zorder=0,
        )
        ax.add_patch(shadow)

        facets = []
        for i in range(cols):
            for j in range(rows):
                quad2 = [grid2[i][j], grid2[i + 1][j], grid2[i + 1][j + 1], grid2[i][j + 1]]
                z_avg = sum(grid3[ii][jj][2] for ii, jj in [(i, j), (i + 1, j), (i + 1, j + 1), (i, j + 1)]) / 4
                depth = sum(grid3[ii][jj][1] for ii, jj in [(i, j), (i + 1, j), (i + 1, j + 1), (i, j + 1)]) / 4
                shade = 0.78 + 0.15 * (z_avg / amp + 1) / 2 - 0.04 * ((i + j) % 2)
                shade = max(0.68, min(0.94, shade))
                facets.append((depth, quad2, shade))

        for _, quad2, shade in sorted(facets, key=lambda item: item[0], reverse=True):
            ax.add_patch(
                Polygon(
                    quad2,
                    closed=True,
                    facecolor=(shade, shade, shade),
                    edgecolor="#2f2f2f",
                    linewidth=0.62,
                    alpha=0.97,
                    joinstyle="round",
                    zorder=1,
                )
            )

        for i in range(cols + 1):
            for j in range(rows):
                p0, p1 = grid2[i][j], grid2[i][j + 1]
                ax.plot([p0[0], p1[0]], [p0[1], p1[1]], color="#1f1f1f", linewidth=0.58, alpha=0.92, zorder=2)
        for j in range(rows + 1):
            for i in range(cols):
                p0, p1 = grid2[i][j], grid2[i + 1][j]
                is_ridge = i % 2 == 0
                ax.plot(
                    [p0[0], p1[0]],
                    [p0[1], p1[1]],
                    color="#111111" if is_ridge else "#4a4a4a",
                    linewidth=0.82 if is_ridge else 0.50,
                    alpha=0.92,
                    zorder=2,
                )
        for i in range(cols):
            for j in range(rows):
                if (i + j) % 2 == 0:
                    p0, p1 = grid2[i][j], grid2[i + 1][j + 1]
                else:
                    p0, p1 = grid2[i + 1][j], grid2[i][j + 1]
                ax.plot([p0[0], p1[0]], [p0[1], p1[1]], color="#3a3a3a", linewidth=0.42, alpha=0.78, zorder=2)

        left_nodes = [grid2[0][j] for j in range(rows + 1)]
        right_nodes = [grid2[cols][j] for j in range(rows + 1)]
        left_nodes.sort(key=lambda point: point[1])
        right_nodes.sort(key=lambda point: point[1])
        return left_nodes, right_nodes

    def add_fixed_boundary(nodes: list[tuple[float, float]], label_y: float) -> None:
        for x, y in nodes:
            ax.add_patch(Circle((x - 0.03, y), 0.065, facecolor=blue, edgecolor=blue, zorder=4))
            ax.plot([x - 0.24, x - 0.24], [y - 0.15, y + 0.15], color=blue, linewidth=1.5)
            for k in range(4):
                ax.plot([x - 0.36, x - 0.24], [y - 0.10 + k * 0.08, y - 0.02 + k * 0.08], color=blue, linewidth=1.0)
        ax.annotate(
            "Fixed left edge\nUx=Uy=Uz=0",
            xy=(nodes[len(nodes) // 2][0] - 0.18, nodes[len(nodes) // 2][1]),
            xytext=(0.55, label_y),
            color=blue,
            fontsize=10.5,
            weight="bold",
            bbox=dict(boxstyle="round,pad=0.32", fc="white", ec=blue, lw=1.2),
            arrowprops=dict(arrowstyle="-", linestyle=(0, (3, 3)), color=blue, lw=1.2),
            ha="left",
            va="center",
        )

    def add_loading(nodes: list[tuple[float, float]], mode: str) -> None:
        for x, y in nodes:
            ax.add_patch(Circle((x + 0.02, y), 0.065, facecolor=red, edgecolor="black", linewidth=0.4, zorder=4))
            if mode == "bending":
                ax.add_patch(FancyArrowPatch((x + 0.10, y - 0.02), (x + 0.10, y - 0.55), arrowstyle="-|>", mutation_scale=15, lw=1.4, color=red))
            else:
                ax.add_patch(FancyArrowPatch((x + 0.12, y), (x + 0.68, y), arrowstyle="-|>", mutation_scale=15, lw=1.4, color=red))

    top_left, top_right = draw_miura_sheet(2.9, 5.15, 0.86)
    bottom_left, bottom_right = draw_miura_sheet(2.9, 2.15, 0.86)
    add_fixed_boundary(top_left, 6.95)
    add_fixed_boundary(bottom_left, 3.70)
    add_loading(top_right, "bending")
    add_loading(bottom_right, "axial")

    ax.plot([0.4, 10.1], [4.55, 4.55], color="#7f7f7f", linestyle=(0, (4, 4)), linewidth=1.0)

    ax.text(
        10.85,
        6.95,
        "Bending:\nforce-control\n-Z load",
        color=red,
        fontsize=11,
        weight="bold",
        ha="center",
        va="center",
        bbox=dict(boxstyle="round,pad=0.32", fc="white", ec=red, lw=1.2),
    )
    ax.text(
        10.85,
        3.25,
        "Axial:\nforce-control\n+X load",
        color=red,
        fontsize=11,
        weight="bold",
        ha="center",
        va="center",
        bbox=dict(boxstyle="round,pad=0.32", fc="white", ec=red, lw=1.2),
    )

    ax.annotate(
        "Crease parameters:\nW, tcrease, creaseE",
        xy=(4.80, 6.15),
        xytext=(6.35, 7.75),
        color=crease_red,
        fontsize=9.4,
        ha="left",
        arrowprops=dict(arrowstyle="->", color=crease_red, lw=1.3),
    )
    ax.annotate(
        "Panel parameters:\ntpanel, panelE",
        xy=(5.75, 2.95),
        xytext=(7.75, 4.05),
        color="#2563eb",
        fontsize=9.4,
        ha="left",
        arrowprops=dict(arrowstyle="->", color="#2563eb", lw=1.3),
    )

    ax.add_patch(FancyArrowPatch((2.65, 1.18), (8.70, 1.18), arrowstyle="<->", mutation_scale=13, color="black", lw=1.2))
    ax.text(5.68, 1.40, "m: unit count / discretization along length", ha="center", va="center", fontsize=8.7)
    ax.add_patch(FancyArrowPatch((2.15, 2.05), (2.15, 3.65), arrowstyle="<->", mutation_scale=13, color="black", lw=1.2))
    ax.text(1.55, 2.86, "n: transverse\nunit count", ha="center", va="center", fontsize=8.7)

    ax.add_patch(FancyArrowPatch((11.00, 5.70), (11.75, 5.70), arrowstyle="-|>", mutation_scale=15, color="black", lw=1.1))
    ax.add_patch(FancyArrowPatch((11.00, 5.70), (11.00, 6.45), arrowstyle="-|>", mutation_scale=15, color=red, lw=1.1))
    ax.add_patch(FancyArrowPatch((11.00, 5.70), (10.58, 5.25), arrowstyle="-|>", mutation_scale=12, color="black", lw=1.0))
    ax.text(11.87, 5.63, "X", fontsize=9)
    ax.text(11.12, 6.48, "Z", fontsize=9)
    ax.text(10.46, 5.16, "Y", fontsize=9)

    ax.text(
        13.45,
        6.10,
        "Simulation settings\n\n"
        "deployment = 0.3 / 0.6 / 0.9\n"
        "80 load steps per curve\n"
        "record force-displacement curve\n"
        "6 curves per sample",
        fontsize=9.2,
        ha="left",
        va="center",
        bbox=dict(boxstyle="round,pad=0.45", fc="white", ec="black", lw=1.0),
    )
    ax.add_patch(FancyArrowPatch((13.30, 5.20), (13.30, 4.65), arrowstyle="simple", mutation_scale=18, color="#9ca3af"))

    x0, y0 = 11.45, 1.45
    width, height = 3.8, 2.25
    ax.add_patch(FancyArrowPatch((x0, y0), (x0 + width, y0), arrowstyle="-|>", mutation_scale=12, lw=1.0, color="black"))
    ax.add_patch(FancyArrowPatch((x0, y0), (x0, y0 + height), arrowstyle="-|>", mutation_scale=12, lw=1.0, color="black"))
    xs = [x0 + width * t / 100 for t in range(101)]
    ys = [
        y0
        + height
        * (0.70 * (t / 100) + 0.16 * (1 / (1 + pow(2.71828, -16 * ((t / 100) - 0.30)))) + 0.08 * max(0, (t / 100) - 0.66))
        for t in range(101)
    ]
    max_y = max(ys)
    ys = [y0 + (y - y0) / (max_y - y0) * (height * 0.95) for y in ys]
    ax.plot(xs, ys, color="#0969da", linewidth=2.0)
    ax.plot([x0, x0 + width * 0.45], [y0, y0 + height * 0.98], color="#ff6b00", linestyle=(0, (4, 3)), linewidth=1.5)
    ax.text(x0 + width / 2, y0 - 0.28, "Displacement", ha="center", va="center", fontsize=9)
    ax.text(x0 - 0.34, y0 + height / 2, "Force", ha="center", va="center", fontsize=9, rotation=90)
    ax.text(x0 + width / 2, y0 + height + 0.18, "Typical nonlinear response", ha="center", va="center", fontsize=9, weight="bold")
    ax.plot([x0 + width * 0.55, x0 + width * 0.72], [y0 + height * 0.78, y0 + height * 0.78], color="#0969da", linewidth=2.0)
    ax.text(x0 + width * 0.75, y0 + height * 0.78, "Nonlinear response", va="center", fontsize=7.2)
    ax.plot(
        [x0 + width * 0.55, x0 + width * 0.72],
        [y0 + height * 0.63, y0 + height * 0.63],
        color="#ff6b00",
        linestyle=(0, (4, 3)),
        linewidth=1.5,
    )
    ax.text(x0 + width * 0.75, y0 + height * 0.63, "Linear stiffness reference", va="center", fontsize=7.2)

    ax.legend(
        handles=[
            Line2D([0], [0], color=blue, lw=2, label="Blue: fixed boundary, all translational DOFs fixed"),
            Line2D([0], [0], color=red, lw=2, label="Red: applied nodal forces at right edge"),
            Rectangle((0, 0), 1, 1, facecolor=panel_fill, edgecolor="#202020", label="Gray: Miura-ori folded sheet, facets and creases"),
        ],
        loc="lower left",
        bbox_to_anchor=(0.02, 0.03),
        fontsize=8.2,
        frameon=False,
    )

    ax.text(
        10.90,
        0.55,
        "SWOMPS NR incremental loading; not top platen compression",
        fontsize=10.5,
        weight="bold",
        ha="center",
        va="center",
        bbox=dict(boxstyle="round,pad=0.35", fc="white", ec="black", lw=1.0, linestyle=(0, (4, 3))),
    )

    fig.tight_layout(pad=0.3)
    fig.savefig(FIG_PARAMETER_SCHEMATIC, bbox_inches="tight")
    plt.close(fig)
    return FIG_PARAMETER_SCHEMATIC


def pct(value: float) -> str:
    return f"{value:.2f}%"


def sci(value: float) -> str:
    if abs(value) >= 1000 or (abs(value) < 0.001 and value != 0):
        return f"{value:.3e}"
    return f"{value:.4g}"


def build_document() -> Path:
    forward = load_json(FORWARD_METRICS)
    inverse = load_json(INVERSE_METRICS)
    swomps = load_json(SWOMPS_SUMMARY)
    audit = load_json(AUDIT_SUMMARY)
    per_candidate = pd.read_csv(SWOMPS_PER_CANDIDATE)

    fwd_curve = forward["denormalized_curve_metrics"]
    inv_curve = inverse["closed_loop_curve_metrics"]
    sw_summary = swomps["summary"]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    parameter_schematic = create_parameter_schematic()
    document = Document()
    set_doc_defaults(document)
    add_page_number(document.sections[0].footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("基于多工况力-位移曲线的 Miura 折纸结构概率逆向设计与 SWOMPS 真实仿真验证")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("小论文初稿 | 作者、单位与基金信息待补充").italic = True

    add_heading(document, "摘要", 1)
    add_para(
        document,
        "针对仅以少量刚度标量作为目标时逆向映射信息不足的问题，本文构建了一个面向 Miura 折纸结构的多工况力-位移曲线数据集，"
        "并建立“原始设计参数到六工况非线性响应曲线”的正向预测模型以及“目标响应曲线到设计参数”的条件扩散逆向模型。"
        "数据集以 SWOMPS 非线性静力加载为仿真核心，输入包括结构离散参数、几何参数与材料参数，输出为弯曲/轴向两类加载在三个部署状态下的 80 点力-位移响应。"
        f"在质量过滤后的随机划分数据集上，正向 GNN-Transformer 模型取得位移 std-NRMSE {pct(fwd_curve['displacement_nrmse_std_percent'])}、"
        f"力 std-NRMSE {pct(fwd_curve['force_nrmse_std_percent'])}。后向 diffusion 在前向代理闭环中取得位移 std-NRMSE "
        f"{pct(inv_curve['displacement_nrmse_std_percent'])}、力 std-NRMSE {pct(inv_curve['force_nrmse_std_percent'])}。"
        f"进一步从生成参数中选取闭环误差最低的 10 个候选设计，重新回到 SWOMPS 进行真实非线性仿真验证，平均位移 std-NRMSE 为 "
        f"{pct(sw_summary['displacement_nrmse_std_percent_mean'])}，平均力 std-NRMSE 为 {pct(sw_summary['force_nrmse_std_percent_mean'])}，"
        "所有样本均收敛且未出现负力或负位移。结果表明，多工况响应曲线相较标量刚度能够提供更丰富的逆向约束，概率生成模型可以在一对多设计空间中产生具有物理可行性的候选结构。"
    )
    add_para(document, "关键词：Miura 折纸；力-位移曲线；逆向设计；扩散模型；GNN-Transformer；SWOMPS；机械超材料")

    add_heading(document, "1 引言", 1)
    add_para(
        document,
        "折纸启发的机械超材料能够通过折痕拓扑、面板厚度、折痕刚度和材料参数调控结构的压缩、弯曲、轴向刚度与能量吸收响应。"
        "在逆向设计任务中，传统数据集常将若干刚度标量作为目标响应，但这种低维物理量只能描述曲线局部斜率，难以表达非线性变形、构型演化、应变能分配以及多工况耦合关系。"
        "因此，仅由刚度到参数的映射容易出现严重的一对多不适定性，模型可能在训练集内表现尚可，却缺乏足够物理信息支撑稳定的生成设计。"
    )
    add_para(
        document,
        "本文工作的出发点是将逆向设计目标从“参数-刚度标量”升级为“参数-多工况响应曲线”。"
        "具体而言，本文基于 Miura-only 数据生成管线，采用大载荷 SWOMPS 非线性响应仿真，为每个样本构造六条 80 点力-位移曲线。"
        "该目标既保留了实际工程中的载荷-变形意义，又比单点刚度包含更多形状、幅值和非线性信息。"
    )

    add_heading(document, "2 数据集与仿真方法", 1)
    add_heading(document, "2.1 参数空间与响应目标", 2)
    add_para(
        document,
        "输入变量为 8 维 Miura 折纸结构参数：pattern、m、n、tcrease、tpanel、W、creaseE 和 panelE。"
        "其中 pattern 固定为 Miura 模式，m 与 n 表示离散结构尺寸，tcrease、tpanel 和 W 描述折痕/面板几何尺度，creaseE 与 panelE 描述折痕和面板材料模量。"
        "输出为 6 个加载工况下的力-位移响应曲线，即弯曲加载和轴向加载分别在初始部署状态 0.3、0.6 和 0.9 下的响应。"
    )
    add_para(
        document,
        "这些变量与折纸结构的对应关系如下。pattern 是结构模式标识，本数据集中固定为 1，表示只研究 Miura-ori 构型；"
        "保留该列是为了与原始数据集格式一致，并为后续混合 Miura/TMP 等多构型实验预留接口。"
        "m 和 n 不是材料参数，而是 Miura 单元离散尺寸参数。"
        "在 MATLAB 几何生成脚本中，纸张基准长度 L 固定为 0.2 m，折痕夹角 gamma 固定为 70 deg，并通过 b=L/m 和 a=(L-b*cot(gamma))/n 计算 Miura 单元边长。"
        "因此，m 或 n 增大时，单元划分更密、局部面板尺度更小，结构整体的变形路径和等效刚度都会随之改变。"
    )
    add_para(
        document,
        "tcrease、tpanel 和 W 是几何-材料耦合参数。"
        "tcrease 表示折痕等效厚度，主要影响折痕转动柔顺性；tpanel 表示面板厚度，主要影响面板拉伸和弯曲抗力；W 表示折痕宽度，同时进入折痕铰链和面板弹簧刚度计算。"
        "在 SWOMPS 的等效杆-铰链模型中，折痕转动刚度近似随 creaseE*tcrease^3/W 增大，面板弯曲相关刚度近似随 panelE*tpanel^3/W 增大。"
        "因此这些参数不是孤立影响某一条曲线，而是共同决定六个加载工况下的整体力-位移响应幅值、曲线斜率和非线性变化程度。"
    )
    add_para(
        document,
        "这里的部署状态 0.3、0.6 和 0.9 指 Miura 折纸在加载开始前的初始展开程度，也就是 MATLAB 几何生成函数 GenerateMiuraSheet 中的 Ext 参数，"
        "并不是载荷比例、压缩比例或应变比例。"
        "在几何生成过程中，Ext 会进入 h=b*sin(gamma)*Ext 和 theta=asin(Ext*sin(gamma)) 等关系式，从而改变节点的三维初始坐标。"
        "因此，deployment=0.3 表示较折叠、起伏更明显且投影长度较短的初始构型；deployment=0.6 表示中等展开状态；deployment=0.9 表示更接近铺展、投影长度更大且起伏较小的初始构型。"
        "同一组结构与材料参数在这三种初始构型下会表现出不同的弯曲和轴向响应，所以每个样本需要分别计算 3 个部署状态乘以 2 种加载方式，共 6 条力-位移曲线。"
    )
    add_para(
        document,
        "这里的“加载工况”指一次具体的仿真条件组合，例如“弯曲加载 + deployment=0.3”就是一个工况。"
        "本文对每个 Miura 样本设置两类加载方式：弯曲加载表示在右侧边界节点施加 -Z 方向力并记录 z 方向位移响应，轴向加载表示在右侧边界节点施加 +X 方向力并记录 x 方向位移响应。"
        "所谓“力-位移响应曲线”则是指在 80 个递增加载步中，逐步记录外力大小和对应结构位移，形成描述结构非线性力学行为的一条曲线。"
        "因此，机器学习模型的输出目标不是单个刚度标量，而是每个样本对应的 6 条离散力-位移曲线；这些曲线共同刻画结构在不同初始展开程度和不同加载方向下的响应。"
    )
    add_table(
        document,
        ["输入变量", "本文取值/范围", "对应结构位置或物理含义", "对响应曲线的主要影响"],
        [
            ["pattern", "1", "结构模式标识；1 表示 Miura-ori", "当前固定，保证 Miura-only 数据集一致性"],
            ["m", "{24, 30, 36}", "沿一个方向的离散尺寸参数；脚本中 b=L/m", "改变单元尺度和折痕数量，影响整体柔顺性与载荷传递路径"],
            ["n", "{6, 9, 12}", "横向/另一方向离散尺寸参数；脚本中 a=(L-b*cot(gamma))/n", "改变单元宽度和面板形状，影响弯曲与轴向响应差异"],
            ["tcrease", "0.0005-0.001 m", "折痕区域等效厚度", "主要控制折痕转动刚度，厚度增大通常使折痕更难转动"],
            ["tpanel", "0.001-0.006 m", "面板区域厚度", "主要控制面板拉伸/弯曲刚度，厚度增大通常提高整体承载能力"],
            ["W", "0.001-0.004 m", "折痕宽度，同时作为面板弹簧计算中的宽度尺度", "影响折痕铰链刚度和面板等效弹簧刚度"],
            ["creaseE", "1.0e9-5.0e9 Pa", "折痕区域等效 Young's modulus", "控制折痕材料刚度，影响折叠/展开阻力"],
            ["panelE", "1.0e9-5.0e9 Pa", "面板区域等效 Young's modulus", "控制面板材料刚度，影响整体力水平和局部变形"],
        ],
        "表 1 输入参数与折纸结构/物理意义对应关系",
    )
    add_table(
        document,
        ["curve_id", "加载类型", "部署状态", "响应方向", "机器学习目标含义"],
        [
            ["0", "bending", "0.3", "z", "较折叠初始构型下的弯曲力-位移曲线"],
            ["1", "bending", "0.6", "z", "中等展开初始构型下的弯曲力-位移曲线"],
            ["2", "bending", "0.9", "z", "较展开初始构型下的弯曲力-位移曲线"],
            ["3", "axial", "0.3", "x", "较折叠初始构型下的轴向力-位移曲线"],
            ["4", "axial", "0.6", "x", "中等展开初始构型下的轴向力-位移曲线"],
            ["5", "axial", "0.9", "x", "较展开初始构型下的轴向力-位移曲线"],
        ],
        "表 2 六个输出工况与曲线编号对应关系",
    )
    add_figure(
        document,
        parameter_schematic,
        "图 1 当前 Miura-ori 的 SWOMPS 仿真加载过程与参数标注。左边界固定全部平动自由度，右边界施加节点力；上图为弯曲工况 -Z 方向力控加载，下图为轴向工况 +X 方向力控加载。m/n 控制单元离散尺寸，tcrease/W/creaseE 对应折痕区域，tpanel/panelE 对应面板区域。",
        width=5.9,
    )
    add_heading(document, "2.2 材料参数化与选择依据", 2)
    add_para(
        document,
        "本研究未将仿真材料限定为某一种具体材料牌号，而是沿用原始 Miura MaterialProperty 数据集的参数化等效材料逻辑。"
        "在原始数据生成脚本中，折痕厚度、面板厚度、折痕宽度、折痕 Young's modulus 和面板 Young's modulus 均作为设计变量随机采样；"
        "因此本文在构造 nonlinear 曲线数据集时保留这些变量，以保持新数据集与原始参数空间的一致性。"
    )
    add_para(
        document,
        "具体而言，折痕与面板被建模为两个等效线弹性区域：tcrease 和 creaseE 控制折痕区域的等效柔顺性，tpanel 和 panelE 控制面板区域的拉伸与弯曲刚度，W 表示折痕宽度。"
        "泊松比固定为 0.3，密度固定为 1200 kg/m3。"
        "由于本文采用的是 SWOMPS Newton-Raphson 准静态加载，力-位移曲线主要由几何构型、厚度、折痕宽度和弹性模量决定，密度不会像动力学仿真中那样主导响应。"
    )
    add_para(
        document,
        "这种选择的优势是能够学习“几何-材料等效参数-多工况响应”的连续映射，并保留折痕与面板力学性质不同这一折纸结构中的关键特征。"
        "如果直接选用 PLA、TPU、铝或树脂等少数具体材料，虽然工程解释更直观，但会压缩材料参数空间，并可能破坏与原数据集的输入定义一致性。"
        "因此，当前数据集更适合作为概率逆向设计模型的主训练集；面向后续论文强化时，可额外选取若干真实材料组合进行 material-specific validation。"
    )
    add_heading(document, "2.3 SWOMPS 非线性仿真流程", 2)
    add_para(
        document,
        "正式数据生成采用 stiffness-scaled 大载荷策略。首先根据小载荷响应估计每条工况曲线的参考刚度，然后令最终载荷等于参考刚度乘以目标线性位移 0.015，并设置最大载荷上限 10000。"
        "每条曲线离散为 80 个加载步，记录 force、displacement、signed_force、signed_displacement、converged 以及应变能、杆件应力/应变、折痕力矩与折痕转角等物理摘要量。"
        "这种设置使数据不再局限于线性刚度复现，而是覆盖更大的几何非线性响应区间。"
    )
    add_heading(document, "2.4 真实验证流程", 2)
    add_para(
        document,
        "为了验证后向模型生成参数是否真正可被 SWOMPS 接受，而不是只在前向代理模型下闭环成立，本文进行了二级验证。"
        "首先从后向 diffusion 输出的 generated_parameters_closed_loop.csv 中按 closed_loop_scaled_mse 选择误差最低的 10 个候选参数。"
        "随后使用这些生成参数重新运行 3 步小载荷 SWOMPS 仿真，按 final_force/final_displacement 估计每个候选结构自己的 6 个参考刚度。"
        "最后使用与原 nonlinear 数据集一致的 80 点大载荷逻辑重新仿真，并将真实 SWOMPS 曲线与对应目标曲线逐点比较。"
    )

    add_heading(document, "3 模型方法", 1)
    add_heading(document, "3.1 正向 GNN-Transformer", 2)
    add_para(
        document,
        "正向模型以 8 维参数为输入，将离散变量通过类别嵌入表示，将连续变量标准化后输入图结构与 Transformer 混合编码器。"
        "输出端采用时序解码结构预测 6 个工况下的 80 点响应曲线，并加入平滑与力单调正则以抑制曲线抖动。"
        "该模型用于快速近似 SWOMPS 响应，并在后向模型候选筛选中作为闭环代理。"
    )
    add_heading(document, "3.2 条件 diffusion 后向模型", 2)
    add_para(
        document,
        "后向模型以目标力-位移曲线为条件，扩散生成 8 维设计参数。"
        "由于折纸逆向设计天然具有一对多特征，单一确定性回归往往难以覆盖所有可行设计；条件 diffusion 通过采样多个候选参数，为每个目标曲线提供一组可能设计。"
        "在候选选择阶段，本文对生成参数进行类别 snap 与连续范围 clip，并引入负物理量惩罚项，使前向代理预测中出现负位移或负力的候选更难被选中。"
    )

    add_heading(document, "4 结果", 1)
    add_table(
        document,
        ["项目", "设置或结果"],
        [
            ["数据集", "data_random_quality_filtered"],
            ["训练/测试样本", f"{forward['metadata']['train_samples']} / {forward['metadata']['test_samples']}"],
            ["响应目标", "6 工况 x 80 步 x 2 通道(displacement, force)"],
            ["材料模型", "折痕/面板分区等效线弹性参数; nu=0.3, rho=1200 kg/m3"],
            ["正向模型", "GNN-Transformer, 50 epochs"],
            ["后向模型", "Conditional diffusion, 150 epochs, 每目标 16 个候选"],
            ["真实验证", "top-10 生成参数, 小载荷刚度估计 + 80 点 nonlinear SWOMPS"],
        ],
        "表 3 数据集、模型与验证配置",
    )
    add_table(
        document,
        ["阶段", "位移 std-NRMSE", "力 std-NRMSE", "补充指标"],
        [
            [
                "正向代理模型",
                pct(fwd_curve["displacement_nrmse_std_percent"]),
                pct(fwd_curve["force_nrmse_std_percent"]),
                f"best epoch {forward['best_epoch']}",
            ],
            [
                "后向模型-前向代理闭环",
                pct(inv_curve["displacement_nrmse_std_percent"]),
                pct(inv_curve["force_nrmse_std_percent"]),
                f"负力点比例 {inv_curve['negative_force_fraction'] * 100:.2f}%",
            ],
            [
                "后向生成参数-真实 SWOMPS 验证(top-10 mean)",
                pct(sw_summary["displacement_nrmse_std_percent_mean"]),
                pct(sw_summary["force_nrmse_std_percent_mean"]),
                "收敛率 100%, 负力/负位移 0",
            ],
            [
                "后向生成参数-真实 SWOMPS 验证(top-10 max)",
                pct(sw_summary["displacement_nrmse_std_percent_max"]),
                pct(sw_summary["force_nrmse_std_percent_max"]),
                "最大误差来自 rank 7",
            ],
        ],
        "表 4 主要预测与验证结果",
    )
    add_figure(document, FIG_FORWARD, "图 2 正向 GNN-Transformer 训练曲线。", width=5.6)
    add_figure(document, FIG_INVERSE, "图 3 后向 diffusion 通过前向代理模型得到的闭环误差。", width=5.6)
    add_figure(document, FIG_SWOMPS_SUMMARY, "图 4 top-10 生成参数回到 SWOMPS 后的真实仿真验证误差。", width=5.8)

    add_table(
        document,
        ["rank", "目标样本", "验证样本", "代理闭环 MSE", "真实位移 std-NRMSE", "真实力 std-NRMSE"],
        [
            [
                str(int(row["rank"])),
                str(row["target_sample_id"]),
                str(row["validation_sample_id"]),
                sci(float(row["closed_loop_scaled_mse"])),
                pct(float(row["displacement_nrmse_std_percent"])),
                pct(float(row["force_nrmse_std_percent"])),
            ]
            for _, row in per_candidate.iterrows()
        ],
        "表 5 top-10 生成参数真实 SWOMPS 验证明细",
    )
    add_figure(document, FIG_SWOMPS_GOOD, "图 5 rank 1 候选设计的真实 SWOMPS 曲线与目标曲线对比。", width=5.8)
    add_figure(document, FIG_SWOMPS_WORST, "图 6 rank 7 候选设计的真实 SWOMPS 曲线与目标曲线对比。", width=5.8)

    add_heading(document, "5 讨论", 1)
    add_heading(document, "5.1 表现好的方面", 2)
    add_para(
        document,
        "第一，真实 SWOMPS 验证结果明显强于仅在前向代理模型中观察到的闭环结果。"
        f"在代理闭环中，力 std-NRMSE 为 {pct(inv_curve['force_nrmse_std_percent'])}，且仍有 {inv_curve['negative_force_fraction'] * 100:.2f}% 的负力预测点；"
        f"而真实 SWOMPS top-10 验证的平均力 std-NRMSE 降至 {pct(sw_summary['force_nrmse_std_percent_mean'])}，并且负力和负位移比例均为 0。"
        "这说明后向模型生成的参数并非只对代理模型有效，至少在 top 候选上能够通过高保真仿真得到物理可行响应。"
    )
    add_para(
        document,
        "第二，曲线目标比单点刚度目标提供了更丰富的设计约束。"
        "真实验证中多数候选在六个工况上均保持良好贴合，rank 1 的弯曲和轴向曲线与目标曲线几乎重合，说明多工况力-位移曲线能够约束结构在不同部署状态下的整体响应形状。"
    )
    add_para(
        document,
        "第三，生成参数的工程合理性较好。"
        "在后向模型候选筛选阶段，pattern、m、n 等类别参数被 snap 到合法集合，连续参数被限制在训练分布范围内。"
        "真实 SWOMPS 审计显示 10 个验证样本均成功输出 480 行曲线数据，issues_count 为 0，说明这些设计没有在几何生成、加载节点识别或 NR 收敛阶段出现明显失败。"
    )

    add_heading(document, "5.2 不足与可能存在的问题", 2)
    add_para(
        document,
        "第一，真实 SWOMPS 验证目前只覆盖 top-10 候选，样本规模仍小。"
        "该验证可以证明模型能够产生一批高质量可行设计，但尚不足以全面说明整个生成分布的稳定性。"
        "若论文需要更强结论，应将验证扩展到 top-50、随机候选和难例候选，并报告不同候选组的成功率和误差分布。"
    )
    add_para(
        document,
        "第二，候选选择存在代理模型偏置。"
        "top-10 是按照前向代理闭环误差筛选的，因此验证结果代表“代理筛选后的优选生成参数”，不等价于 diffusion 原始采样分布的无偏性能。"
        "后续应补充未筛选样本、不同采样温度、不同候选数 K 的消融实验，以区分后向生成能力和代理筛选能力。"
    )
    add_para(
        document,
        "第三，前向代理模型仍存在低力区域物理边界问题。"
        "代理闭环结果中约 3.80% 的力点为负，尽管真实 SWOMPS 验证中未出现负力，但这说明代理模型在低力起始段的物理一致性仍有改进空间。"
        "可以考虑在正向模型中采用非负力增量累积解码、边界条件约束或更强的低载荷局部损失。"
    )
    add_para(
        document,
        "第四，真实验证采用同源 SWOMPS 仿真，而不是实验样机或其他有限元软件交叉验证。"
        "因此当前结论主要说明模型学习到了 SWOMPS 数据生成逻辑下的有效逆向规律，尚不能直接推广为真实物理样机性能。"
        "若面向高水平论文投稿，建议增加 Abaqus/实验压缩测试或至少增加不同仿真参数下的稳健性验证。"
    )
    add_para(
        document,
        "第五，目前的 nonlinear 响应未出现 snap-like 或位移非单调响应。"
        f"审计结果显示 snap_like_curves 为 {audit['snap_like_curves']}，nonmonotonic_segments_total 为 {audit['nonmonotonic_segments_total']}。"
        "这意味着当前大载荷范围已包含明显的非线性硬化/软化信息，但尚未覆盖跳跃、屈曲后路径或多稳态分支。"
        "如果论文希望强调多稳态和屈曲设计，需要引入 MGDCM/DC 延拓或更强加载路径。"
    )

    add_heading(document, "6 结论", 1)
    add_para(
        document,
        "本文基于 Miura 折纸结构构建了多工况力-位移曲线数据集，并将其用于正向代理建模和条件 diffusion 逆向生成。"
        "相较于参数-刚度标量数据，曲线数据提供了更高维、更具工程意义的响应约束，使后向模型能够在一对多设计空间中产生物理可行的候选参数。"
        f"真实 SWOMPS top-10 验证表明，生成设计的平均位移 std-NRMSE 为 {pct(sw_summary['displacement_nrmse_std_percent_mean'])}，"
        f"平均力 std-NRMSE 为 {pct(sw_summary['force_nrmse_std_percent_mean'])}，且所有验证样本均收敛。"
        "这些结果支持将“多工况响应曲线-概率生成模型-高保真仿真闭环验证”作为折纸结构逆向设计的一条可行路线。"
    )

    add_heading(document, "7 后续工作计划", 1)
    add_para(
        document,
        "后续建议按三个层次推进：其一，扩大真实 SWOMPS 验证规模，至少覆盖 top-50、随机 50 和难例 50；其二，补充模型消融，包括候选数 K、物理惩罚权重、非负增量解码、物理摘要多任务学习等；"
        "其三，引入跨仿真或实验验证，检查生成参数在不同求解器、不同载荷路径和真实样机中的一致性。"
    )

    add_heading(document, "数据与文件位置", 1)
    add_para(document, f"质量过滤数据集：{EXPERIMENT_DIR / 'data_random_quality_filtered'}")
    add_para(document, f"后向 diffusion 结果：{INVERSE_METRICS.parent}")
    add_para(document, f"真实 SWOMPS 验证目录：{VALIDATION_DIR}")
    add_para(document, f"逐候选指标：{SWOMPS_PER_CANDIDATE}")
    add_para(document, f"审计结果：{AUDIT_SUMMARY}")

    add_heading(document, "参考文献待补", 1)
    add_para(
        document,
        "本初稿尚未补充正式参考文献。后续至少应补充 Miura-ori 结构力学、SWOMPS/OrigamiSimulator、GraphMetaMat 或相关图神经网络材料逆向设计、扩散模型逆向设计等方向的文献。"
    )

    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_document()
    print(path)
